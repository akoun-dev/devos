"""Task 35 — TASKS.xlsx : insertion TASK-020 (vues w-full web + navigation mobile tiroir seul) + resynchronisation.
Motif : copies de styles ligne 23 (TASK-019), plages STATISTIQUES étendues à 5:24, TOTAL 20 tâches.
"""
from copy import copy

from openpyxl import load_workbook

PATH = "/home/z/my-project/.ai/TASKS.xlsx"

wb = load_workbook(PATH)
ws = wb["TASKS"]

# 1. Localiser la ligne TOTAL
total_row = None
for r in range(5, ws.max_row + 1):
    if ws.cell(row=r, column=2).value == "TOTAL":
        total_row = r
        break
assert total_row == 24, f"TOTAL attendu en 24, trouvé {total_row}"

heights = {
    r: ws.row_dimensions[r].height
    for r in list(ws.row_dimensions.keys())
    if ws.row_dimensions[r].height is not None
}

# 2. Insérer la ligne TASK-020 juste avant TOTAL (styles copiés de TASK-019, ligne 23)
ws.insert_rows(total_row)
new_row = total_row  # 24
src_row = 23
for c in range(2, 21):
    src = ws.cell(row=src_row, column=c)
    dst = ws.cell(row=new_row, column=c)
    dst._style = copy(src._style)

# restaurer les hauteurs décalées
for old_r, h in heights.items():
    target = old_r + 1 if old_r >= total_row else old_r
    ws.row_dimensions[target].height = h

values = {
    2: "TASK-020",
    3: "PRODUIT",
    4: "DevOS — Espace de travail (FEATURE-UI-SHELL)",
    5: "Vues en pleine largeur (w-full) sur web + navigation mobile réduite au tiroir",
    6: (
        "Exigence CTO « Mets les vues en w-full pour web, en mobile, retire le menu compacte, garde juste "
        "la sidebar (drawer) ». (1) main passe de mx-auto max-w-5xl à w-full (preuves mesurées : 1280 → main "
        "1024 px = 1280−256 sidebar ; 1920 → main 1664 px, max-width none, marges 0) ; footer aligné w-full. "
        "(2) BottomNav (barre inférieure 5 destinations) supprimée : composant bottom-nav.tsx supprimé, champ "
        "mobilePrimary + export MOBILE_PRIMARY_ITEMS retirés du registre navigation.ts, dégagements retirés "
        "(pb-24 sur main, +4,5rem sur footer). (3) Navigation mobile = tiroir (Sheet) seul via burger header : "
        "E2E 375 px — bottomNavPresent=false, burger visible, tiroir complet (11 actives + 5 « Bientôt »), "
        "navigation Projets OK avec fermeture auto. ADR-004 amendé (§4). Aucun changement de contrat API."
    ),
    7: "Frontend / UX (CTO Task 35)",
    8: "TERMINÉ",
    9: 1,
    10: "P2",
    11: "TASK-019 (FEATURE-UI-SHELL — le shell et la navigation existent)",
    12: (
        "src/components/shell/app-shell.tsx; src/lib/navigation.ts; src/components/shell/bottom-nav.tsx "
        "(supprimé); .ai/ADR/ADR-004-shell-navigation-permissions.md; tool-results/task20-desktop-wfull.png; "
        "tool-results/task20-desktop-1920-wfull.png; tool-results/task20-mobile-drawer.png; "
        "tool-results/task20-mobile-projets.png; worklog.md"
    ),
    13: (
        "Desktop 1280 : mainW=1024 (= viewport−sidebar 256), mainMaxW=none, marges 0 · Desktop 1920 : "
        "mainW=1664, footerW=1664, footerMaxW=none · Mobile 375 : bottomNavPresent=false, burger présent+visible, "
        "mainW=375, hasHScroll=false · E2E tiroir : ouverture → navigation complète visible → clic Projets → "
        "h1 « Projets » + tiroir fermé · lint 0 · tsc 0 · 0 erreur console (2 warnings Radix DialogContent "
        "préexistants, hors périmètre)"
    ),
    14: "PASS",
    15: 0,
    16: (
        "0 middleware respecté (ADR-003) · registre à source unique conservé (sidebar + tiroir générés depuis "
        "NAV_SECTIONS — aucune duplication) · a11y : navigation tiroir aria-labelée, aria-current conservé, "
        "cibles ≥ 44 px (min-h-11) · suppression fichier = dette code réduite (−76 l. bottom-nav.tsx)"
    ),
    17: (
        "Web : vues exploitent toute la largeur disponible (fini la colonne max-w-5xl) · Mobile : navigation "
        "unifiée tiroir (une seule source de vérité visuelle), moins de chrome à l'écran · MOBILE_PRIMARY_ITEMS "
        "retiré — breaking change interne sans impact externe (aucun autre consommateur)"
    ),
    18: "2026-10-26",
    19: "2026-10-26",
    20: (
        "AUDIT-006 requis après cette évolution (règle « audit après chaque feature ») · le cahier des charges "
        "initial §5 (barre inférieure 5 destinations) est officiellement supplanté par l'amendement ADR-004 §4"
    ),
}
for c, v in values.items():
    ws.cell(row=new_row, column=c).value = v

# 3. Ligne TOTAL (décalée en 25)
tr = total_row + 1
ws.cell(row=tr, column=3).value = "20 tâches"
ws.cell(row=tr, column=9).value = 1
ws.cell(row=tr, column=15).value = 0
ws.cell(row=tr, column=20).value = "Terminées : 20 · En cours : 0 · À faire : 0 · Bloquées : 0"

# 4. STATISTIQUES — plages étendues aux lignes 5..24 (20 tâches)
st = wb["STATISTIQUES"]
st["C5"] = 20
st["C6"] = '=COUNTIF(TASKS!H5:H24,"TERMINÉ")'
st["C7"] = '=COUNTIF(TASKS!H5:H24,"EN COURS")'
st["C8"] = '=COUNTIF(TASKS!H5:H24,"À FAIRE")'
st["C9"] = '=COUNTIF(TASKS!H5:H24,"BLOQUÉ")'
st["C10"] = "=SUM(TASKS!O5:O24)"
st["C11"] = "=AVERAGE(TASKS!I5:I24)"

wb.save(PATH)
print("TASKS.xlsx mis à jour : TASK-020 insérée en ligne 24, TOTAL en 25, STATISTIQUES resynchronisées (5:24).")
