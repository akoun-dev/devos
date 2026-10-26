"""Task 38 / TASK-024 — TASKS.xlsx : refonte UI « DevOS Intelligence » selon maquette CTO.
Motif : copie de styles ligne 27 (TASK-023), insertion 1 ligne avant TOTAL (28), TOTAL → 29, 24 tâches.
"""
from copy import copy

from openpyxl import load_workbook

PATH = "/home/z/my-project/.ai/TASKS.xlsx"

wb = load_workbook(PATH)
ws = wb["TASKS"]

total_row = None
for r in range(5, ws.max_row + 1):
    if ws.cell(row=r, column=2).value == "TOTAL":
        total_row = r
        break
assert total_row == 28, f"TOTAL attendu en 28, trouvé {total_row}"

heights = {
    r: ws.row_dimensions[r].height
    for r in list(ws.row_dimensions.keys())
    if ws.row_dimensions[r].height is not None
}

# Insertion d'1 ligne avant TOTAL (styles copiés de TASK-023, ligne 27)
ws.insert_rows(total_row, amount=1)
src_row = 27
for c in range(2, 21):
    src = ws.cell(row=src_row, column=c)
    dst = ws.cell(row=total_row, column=c)
    dst._style = copy(src._style)

for old_r, h in heights.items():
    target = old_r + 1 if old_r >= total_row else old_r
    ws.row_dimensions[target].height = h

# ─── TASK-024 — Refonte UI « DevOS Intelligence » ───
values_24 = {
    2: "TASK-024",
    3: "PRODUIT",
    4: "DevOS — Refonte UI Assistant (maquette CTO)",
    5: "Vue « DevOS Intelligence » refondue : badge Unifie, bienvenue, composer voix/texte, chips",
    6: (
        "Exigence CTO « pour la vue DEVOS INTELLIGENCE, je veux ceci » (maquette 1). ZÉRO changement backend "
        "(contrat v1.4 inchangé champ à champ). UI : en-tête logo carré bleu BrainCircuit + titre « DevOS "
        "Intelligence » + badge vert « Unifie » + sous-titre « Assistant IA unifié · Texte et voix » ; carte "
        "de bienvenue (état vide) « Bonjour ! Je suis votre assistant DevOS. Comment puis-je vous aider ? » "
        "+ « Vous pouvez écrire votre demande ou utiliser le microphone. » + suggestions contextuelles ; "
        "composer : placeholder « Posez votre question ou donnez une commande... », bouton « Pièce jointe » "
        "(toast honnête « bientôt » — ADR-004 §5, aucun faux lien mort), bouton micro (dictée Web Speech via "
        "NOUVEAU hook use-speech-dictation fr-FR — transcript inséré dans le champ POUR relecture avant "
        "envoi, dégradation gracieuse toast si API absente), bouton envoi flèche ↑ ; fil de messages quand "
        "conversation active (bulles arrondies, badge « Persistée », quota restant en toolbar) ; lecture "
        "vocale des réponses (bouton Volume2 → lib/assistant/tts) ; chips de capacités : « Historique "
        "partagé · N » (OUVRE un Sheet latéral avec la liste des conversations persistées : reprise, "
        "suppression, compteur dynamique), « Réponses vocales », « Commandes DevOS » ; « Nouvelle "
        "conversation » conservé ; compteur 0/4000 progressif (> 3500). Fonctionnalités TASK-021 intégralement "
        "préservées (reprise/suppression historique, rate limit, 401)."
    ),
    7: "Équipe complète (CTO Task 38)",
    8: "TERMINÉ",
    9: 1,
    10: "P1",
    11: "TASK-021 (historique IA persistant — Sheet) + TASK-019 (vue shell)",
    12: (
        "src/components/views/ai-chat-view.tsx; src/hooks/use-speech-dictation.ts; "
        "tool-results/task24-desktop-bienvenue.png; tool-results/task24-desktop-conversation.png; "
        "tool-results/task24-historique-sheet.png; tool-results/task24-mobile-bienvenue.png; "
        "tool-results/task24-mobile-conversation.png; worklog.md"
    ),
    13: (
        "E2E agent-browser : desktop 1280 rendu maquette (en-tête + bienvenue + composer + chips, capture) · "
        "toast « Pièces jointes — bientôt disponible » prouvé · dictée : SpeechRecognition mocké → transcript "
        "« bonjour depuis la diction mockee » inséré dans le champ (preuve get value) · envoi réel → réponse "
        "IA contextuelle (2 projets actifs cités) + quota 9 affiché + bouton « Écouter cette réponse » · "
        "Sheet « Historique des conversations » : liste 3 conv, reprise « salut » (2 réponses restaurées), "
        "suppression ×2 de MES artefacts de test (purge prouvée : chip 3 → 1, « salut » préexistante "
        "préservée) · « Nouvelle conversation » → retour bienvenue texte maquette exact · mobile 375 : tiroir "
        "→ vue w=375 hasHScroll=false, cycle complet + captures · 0 erreur console · lint 0 · tsc 0"
    ),
    14: "PASS",
    15: 0,
    16: (
        "0 endpoint touché (aucune régression API possible) · 0 secret · SDK serveur uniquement · dictée "
        "client-only (aucun audio transmis) · TTS via Web Speech local (lib/assistant/tts) · a11y : "
        "aria-pressed/aria-haspopup/aria-label sur tous les contrôles, cibles ≥ 44px, SheetTitle+Description "
        "(pas de warning Radix) · purge des artefacts de test effectuée"
    ),
    17: (
        "La vue Assistant devient l'écran « vitrine » de la maquette CTO : identité « DevOS Intelligence » + "
        "capacités visibles (historique, voix, commandes) sans régression fonctionnelle — l'historique "
        "persistant passe du panneau statique au Sheet, plus épuré, fidèle au cahier des charges"
    ),
    18: "2026-10-26",
    19: "2026-10-26",
    20: (
        "Pièce jointe réelle (upload + contrat v1.5) candidate backlog · STT serveur (transcription "
        "z-ai-web-dev-sdk) si Web Speech insuffisant · refonte du floating-assistant sur le hook partagé "
        "(duplication SpeechRecognition acceptée cette session — 2 sites seulement)"
    ),
}
for c, v in values_24.items():
    ws.cell(row=total_row, column=c).value = v

# TOTAL décalé en 29
tr = total_row + 1
ws.cell(row=tr, column=3).value = "24 tâches"
ws.cell(row=tr, column=9).value = 1
ws.cell(row=tr, column=15).value = 0
ws.cell(row=tr, column=20).value = "Terminées : 24 · En cours : 0 · À faire : 0 · Bloquées : 0"

# STATISTIQUES — plages 5..28 (24 tâches)
st = wb["STATISTIQUES"]
st["C5"] = 24
st["C6"] = '=COUNTIF(TASKS!H5:H28,"TERMINÉ")'
st["C7"] = '=COUNTIF(TASKS!H5:H28,"EN COURS")'
st["C8"] = '=COUNTIF(TASKS!H5:H28,"À FAIRE")'
st["C9"] = '=COUNTIF(TASKS!H5:H28,"BLOQUÉ")'
st["C10"] = "=SUM(TASKS!O5:O28)"
st["C11"] = "=AVERAGE(TASKS!I5:I28)"

wb.save(PATH)
print("TASKS.xlsx : TASK-024 insérée (28), TOTAL en 29 = 24 tâches, STATISTIQUES 5:28.")
