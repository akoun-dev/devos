"""Task 37 — TASKS.xlsx : insertion TASK-021/022/023 (v1.4 : historique IA, automatisations, solde M3).
Motif : copies de styles ligne 24 (TASK-020), plages STATISTIQUES étendues à 5:27, TOTAL 23 tâches.
"""
from copy import copy

from openpyxl import load_workbook

PATH = "/home/z/my-project/.ai/TASKS.xlsx"

wb = load_workbook(PATH)
ws = wb["TASKS"]

total_row = None
for r in range(5, ws.max_row + 1):
    if ws.cell(row=r, column=2).value == "TOTAL":
        total_row = r
        break
assert total_row == 25, f"TOTAL attendu en 25, trouvé {total_row}"

heights = {
    r: ws.row_dimensions[r].height
    for r in list(ws.row_dimensions.keys())
    if ws.row_dimensions[r].height is not None
}

# Insertion de 3 lignes avant TOTAL (styles copiés de TASK-020, ligne 24)
ws.insert_rows(total_row, amount=3)
src_row = 24
for offset in range(3):
    new_row = total_row + offset
    for c in range(2, 21):
        src = ws.cell(row=src_row, column=c)
        dst = ws.cell(row=new_row, column=c)
        dst._style = copy(src._style)

for old_r, h in heights.items():
    target = old_r + 3 if old_r >= total_row else old_r
    ws.row_dimensions[target].height = h

# ─── TASK-021 — Historique IA persistant ───
values_21 = {
    2: "TASK-021",
    3: "PRODUIT",
    4: "DevOS — Itérations v1.4 (FEATURE-AI-HISTORY)",
    5: "Historique IA persistant (conversations en base + UI de reprise)",
    6: (
        "Exigence CTO « enchaîne sur les itérations suivantes (historique IA persistant, …) ». Contrat v1.4 "
        "écrit AVANT le code. +2 modèles Prisma (AiConversation/AiMessage, cascade user, index userId+updatedAt "
        "et conversationId+createdAt) ; 5 endpoints /api/ai/conversations (liste limit 1..50 tri updatedAt desc, "
        "création, détail+messages asc, renommage title 1..160, suppression cascade) ; POST /api/ai/chat étendu "
        "(conversationId optionnel vérifié → 404 générique ; absent → création implicite avec titre auto = "
        "1ᵉʳ message user tronqué 60 ; persiste dernier message user + réponse — l'historique n'est JAMAIS "
        "réinjecté au LLM au-delà des 20 messages fournis) ; UI : liste « Conversations récentes » (reprise, "
        "suppression, compteur messages, badge Historique enregistré). Rate limit chat inchangé 10/5 min/IP."
    ),
    7: "Équipe complète (CTO Task 37)",
    8: "TERMINÉ",
    9: 1,
    10: "P1",
    11: "TASK-019 (FEATURE-UI-SHELL — la vue Assistant IA existait)",
    12: (
        "prisma/schema.prisma; src/lib/serialize.ts; src/app/api/ai/chat/route.ts; "
        "src/app/api/ai/conversations/route.ts; src/app/api/ai/conversations/[id]/route.ts; "
        "src/components/views/ai-chat-view.tsx; src/types/workspace.ts; .ai/API_CONTRACTS.md (v1.4); "
        "tool-results/task37-ai-histoire-restauree.png; worklog.md"
    ),
    13: (
        "curl : chat sans conversationId → 200 + conversationId, conversation auto-titrée « Donne-moi un conseil… » "
        "messageCount 2 · chat #2 avec conversationId → 200 (LLM réel, prénom injecté) · rename 200 · title vide "
        "→ 400 Zod · inter-users demo → 404 générique · sans cookie → 401 · E2E UI : envoi message → 2 bulles + "
        "conversation listée avec titre auto · reload complet → conversation conservée → clic → 2 messages "
        "restaurés (capture) · purge UI : suppression conversation → « Aucune conversation enregistrée » · "
        "lint 0 · tsc 0"
    ),
    14: "PASS",
    15: 0,
    16: (
        "Scopage userId strict (findFirst {id,userId} → 404 générique) · cascade delete vérifiée (messages) · "
        "cap 20 messages au LLM conservé (aucune fuite d'historique long) · rate limit intact · SDK serveur "
        "uniquement · aucune donnée privée injectée au prompt (indice de comptes inchangé)"
    ),
    17: (
        "Les conversations survivent au rechargement et à la déconnexion — le backlog « historique IA » de "
        "Task 33 est soldé · reprise multi-sessions possible · UI honnête (compteur de messages, dates)"
    ),
    18: "2026-10-26",
    19: "2026-10-26",
    20: (
        "Renommage manuel (PATCH) disponible via API, non exposé en UI v1 (titre auto suffisant) — candidat "
        "backlog · limit de l'historique 50 conversations"
    ),
}
for c, v in values_21.items():
    ws.cell(row=25, column=c).value = v

# ─── TASK-022 — Automatisations ───
values_22 = {
    2: "TASK-022",
    3: "PRODUIT",
    4: "DevOS — Itérations v1.4 (FEATURE-AUTOMATIONS)",
    5: "Automatisations : scénarios horaires et conditions, exécution par le tick + run manuel",
    6: (
        "Exigence CTO « …automatisations… ». Contrat v1.4 AVANT le code. +2 modèles Prisma (Automation/"
        "AutomationRun, cascade, @@unique(automationId,scheduledFor) = registre de dédup DÉDIÉ — le cœur du "
        "scheduler ADR-002 reste INCHANGÉ : taskId null ne peut pas réutiliser ReminderLog dont la contrainte "
        "traiterait les NULL comme distincts en SQLite) ; module src/lib/assistant/automations.ts (pur "
        "evaluateAutomations appelé APRÈS runTick dans le tick ; runAutomationManuellement idempotent à la "
        "seconde) ; 2 déclencheurs : « schedule » (HH:MM UTC + jours CSV 0..6 + message personnalisé 0..500, "
        "grâce 15 min, 1 exécution/jour/slot) et « condition » overdue_count (tâches en_cours avec dueAt < now "
        "≥ seuil 1..99, 1 exécution/jour max via bucket startOfUtcDay) ; messages TEMPLATE uniquement (aucune "
        "invention LLM), jamais urgent, silencieux → run « expire » ; 4 endpoints /api/automations (CRUD avec "
        "validation croisée Zod superRefine + PATCH partiel validé sur l'état FUTUR + run manuel 409 si "
        "double-clic < 1 s, catch P2002 exclusivement — autres erreurs DB propagées) ; vue ACTIVE (registre "
        "planned→active, matrice permissions déjà prévue) : cartes (badge type, résumé déclencheur, switch "
        "enabled, dernière exécution), dialog création/édition (jours checkboxes, heure, message), test "
        "manuel, historique 10 runs, suppression confirmée ; badge annonce « Automatisation » dans le journal "
        "(assistant-panel TYPE_CONFIG étendu)."
    ),
    7: "Équipe complète (CTO Task 37)",
    8: "TERMINÉ",
    9: 1,
    10: "P1",
    11: "TASK-019 (vue planned « Bientôt » préexistante) + FEATURE-001 (pipeline annonces)",
    12: (
        "src/lib/assistant/automations.ts; src/app/api/automations/route.ts; "
        "src/app/api/automations/[id]/route.ts; src/app/api/automations/[id]/run/route.ts; "
        "src/app/api/assistant/tick/route.ts; src/lib/assistant/templates.ts; "
        "src/components/views/automations-view.tsx; src/components/assistant/dev-os-app.tsx; "
        "src/lib/navigation.ts; src/hooks/use-workspace.ts; tool-results/task37-automatisation-creee.png; "
        "tool-results/task37-automatisation-historique.png; tool-results/task37-automatisations-mobile.png"
    ),
    13: (
        "curl : POST schedule → 201 · condition sans threshold → 400 (validation croisée) · run #1 → 200 "
        "(trigger manuel, message personnalisé, firedAt) · run #2 même seconde → 409 · détail + runs 200 · "
        "PATCH enabled:false 200 · E2E tick fort : tâche en retard réelle + condition automation → tick#1 "
        "annonce type 'automation' parlée (« 1 tâche en retard (seuil : 1) ») · tick#2 → [] DÉDUP PROUVEE · "
        "lastRunAt + run condition « annonce » persistés · E2E UI : création dialog (Lun-Ven, 08:00, message) "
        "→ carte + badge + switch · Tester → toast + run manuel · Historique → run visible (capture) · "
        "suppression confirmée → « Aucune automatisation » · mobile 375 : tiroir → vue fonctionnelle, "
        "hasHScroll=false · 0 erreur console · purge artefacts : 0 automatisation / 0 conversation / "
        "0 tâche test · lint 0 · tsc 0"
    ),
    14: "PASS",
    15: 0,
    16: (
        "Cœur ADR-002 non modifié (module séparé, fusion en fin de tick) · horaires documentés UTC (fuseau "
        "serveur Africa/Abidjan) · jamais urgent → jamais de contournement des heures silencieuses · "
        "idempotence par contrainte DB + résolution préalable · 404 générique inter-users · Zod strict avec "
        "validation croisée superRefine · a11y : labels, switch aria-labelé, dialogs décrits"
    ),
    17: (
        "Le menu « Automatisations — Bientôt » devient une fonctionnalité réelle (1 des 5 menus avancés "
        "activés sans bouleverser la navigation — recommandation §6 du cahier des charges respectée) · "
        "l'assistant vocal annonce les scénarios via le pipeline existant (journal + TTS + notifications)"
    ),
    18: "2026-10-26",
    19: "2026-10-26",
    20: (
        "MVP v1 : messages template uniquement (pas de résumé IA composé au slot) ; conditions limitées à "
        "overdue_count — autres conditions (nb de tâches terminées, streaks) candidates backlog · horaires UTC "
        "(conversion locale client non faite)"
    ),
}
for c, v in values_22.items():
    ws.cell(row=26, column=c).value = v

# ─── TASK-023 — Solde des 3 drifts documentaires M3 ───
values_23 = {
    2: "TASK-023",
    3: "GOUVERNANCE",
    4: "DevOS — Solde M3 hérité (AUDIT-005/006)",
    5: "Solde des 3 drifts documentaires : TEAM_STATUS.md, ARCHITECTURE.md §4, CHANGELOG.md",
    6: (
        "Constat M3 récurrent AUDIT-005/006 (3 sessions). (1) TEAM_STATUS.md réaligné intégralement : en-tête "
        "« AUDIT-002 PENDING » → état réel (AUDIT-006 81,3/100 AUTORISÉ, AUDIT-007 en cours), KPI 14→23 tâches, "
        "ligne SECURITY « SAMEORIGIN » → XFO DENY + CSP 10 directives, répartition équipe réécrite. "
        "(2) ARCHITECTURE.md §4 : modèle de données « démo scaffold » (User/Post + warning relation non "
        "déclarée) → état réel 10 tables avec relations et index (Post: relation déclarée, warning retiré ; "
        "note DEBT-001 conservée). (3) CHANGELOG.md : « 9 directives » → 10 (correction M2 AUDIT-004) + "
        "entrée [0.6.0-workspace-iterations] ajoutée (TASK-019/020/021/022 — les versions 0.6 n'avaient "
        "jamais été consignées)."
    ),
    7: "DOC (CTO Task 37)",
    8: "TERMINÉ",
    9: 1,
    10: "P2",
    11: "AUDIT-006 M3 (plan d'action)",
    12: ".ai/TEAM_STATUS.md; .ai/ARCHITECTURE.md; .ai/CHANGELOG.md; worklog.md",
    13: (
        "rg verrou : 0 occurrence résiduelle de « AUDIT-002 PENDING », « SAMEORIGIN » et « 9 directives » dans "
        "les 3 fichiers · cohérence croisée TASKS.md 23/23 · les faits cités (scores, dates, commits) "
        "recoupés avec AUDIT_REPORT.md et COMMIT_LOG.md"
    ),
    14: "PASS",
    15: 0,
    16: "Aucun fichier applicatif touché · correction factuelle uniquement (aucune invention)",
    17: "M3 soldé après 3 sessions récurrentes — la dette documentaire ne masque plus l'état réel du projet",
    18: "2026-10-26",
    19: "2026-10-26",
    20: "Vigilance : les rapports d'équipe doivent être mis à jour À CHAQUE feature (règle rappelée au worklog)",
}
for c, v in values_23.items():
    ws.cell(row=27, column=c).value = v

# TOTAL décalé en 28
tr = total_row + 3
ws.cell(row=tr, column=3).value = "23 tâches"
ws.cell(row=tr, column=9).value = 1
ws.cell(row=tr, column=15).value = 0
ws.cell(row=tr, column=20).value = "Terminées : 23 · En cours : 0 · À faire : 0 · Bloquées : 0"

# STATISTIQUES — plages 5..27 (23 tâches)
st = wb["STATISTIQUES"]
st["C5"] = 23
st["C6"] = '=COUNTIF(TASKS!H5:H27,"TERMINÉ")'
st["C7"] = '=COUNTIF(TASKS!H5:H27,"EN COURS")'
st["C8"] = '=COUNTIF(TASKS!H5:H27,"À FAIRE")'
st["C9"] = '=COUNTIF(TASKS!H5:H27,"BLOQUÉ")'
st["C10"] = "=SUM(TASKS!O5:O27)"
st["C11"] = "=AVERAGE(TASKS!I5:I27)"

wb.save(PATH)
print("TASKS.xlsx : TASK-021/022/023 insérées (25-27), TOTAL en 28 = 23 tâches, STATISTIQUES 5:27.")
