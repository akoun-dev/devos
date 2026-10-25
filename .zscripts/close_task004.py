#!/usr/bin/env python3
"""Mise à jour de TASKS.xlsx : clôture TASK-004 (AUDIT-001 terminé, score 70/100)."""
import sys, os

XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

OUT = "/home/z/my-project/.ai/TASKS.xlsx"
wb = load_workbook(OUT)
ws = wb["TASKS"]

# Ligne TASK-004 = row 8 (row 4 = headers, data starts row 5 : 001→5, 002→6, 003→7, 004→8)
row = 8
assert ws.cell(row=row, column=2).value == "TASK-004", f"Attendu TASK-004, trouvé {ws.cell(row=row, column=2).value}"

ACCENT_POSITIVE = "1B7D46"
ws.cell(row=row, column=8, value="TERMINÉ")            # Statut
ws.cell(row=row, column=8).font = Font(size=10, bold=True, color=ACCENT_POSITIVE)
ws.cell(row=row, column=8).fill = PatternFill("solid", start_color="E8F5E9", end_color="E8F5E9")
ws.cell(row=row, column=9, value=1.0)                   # Progression
ws.cell(row=row, column=9).number_format = "0%"
ws.cell(row=row, column=14, value="PASS")               # Résultat test
ws.cell(row=row, column=19, value="2026-10-25")         # Date fin
ws.cell(row=row, column=20, value="AUDIT-001 publié : score global 70/100 · 0 CRITIQUE · 15 items de dette (7 MAJEUR/8 MINEUR) · Conformité système 58% · Verdict PROD : AUTORISÉ sous condition (AUDIT-002 avant 1re mise en prod réelle)")

# Ligne TOTAL (row 10) — corrige aussi le bug génération : le label TOTAL était écrasé en col B
total_row = 10
assert ws.cell(row=total_row, column=9).value is not None
ws.cell(row=total_row, column=2, value="TOTAL")
ws.cell(row=total_row, column=3, value="5 tâches")
ws.cell(row=total_row, column=9, value=0.8)             # progression moyenne (1.0×4 + 0.0)/5
ws.cell(row=total_row, column=9).number_format = "0%"
ws.cell(row=total_row, column=20, value="Terminées : 4 · En cours : 0 · À faire : 1 · Bloquées : 0")

# Feuille STATISTIQUES : corrige format % sur la bonne ligne (11 = Progression moyenne)
ws2 = wb["STATISTIQUES"]
ws2.cell(row=11, column=3).number_format = "0%"
ws2.cell(row=12, column=3, value="70 / 100 (AUDIT-001, 2026-10-25)")
ws2.cell(row=12, column=3).font = Font(size=10, bold=True, color=ACCENT_POSITIVE)
ws2.cell(row=13, column=3, value=15)
ws2.cell(row=13, column=4, value=".ai/DEBT_REPORT.md (7 MAJEUR / 8 MINEUR / 0 CRITIQUE)")

wb.save(OUT)

# Vérification programmatique
wb2 = load_workbook(OUT)
w = wb2["TASKS"]
print("TASK-004 statut :", w.cell(row=8, column=8).value, "| prog :", w.cell(row=8, column=9).value)
print("TOTAL prog :", w.cell(row=10, column=9).value)
print("STATISTIQUES santé :", wb2["STATISTIQUES"].cell(row=12, column=3).value)
print("OK")
