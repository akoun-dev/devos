"""Task 28 — TASKS.xlsx : insertion TASK-017 (durcissement sécurité + REV-004) + resynchronisation TOTAL/STATISTIQUES.
Motif : copies de styles ligne 20 (TASK-016), plages STATISTIQUES étendues à H5:H21 (17 tâches).
"""
from copy import copy

from openpyxl import load_workbook

PATH = "/home/z/my-project/.ai/TASKS.xlsx"

wb = load_workbook(PATH)
ws = wb["TASKS"]

total_row = None
for r in range(5, ws.max_row + 1):
    if ws.cell(row=r, column=2).value == "TOTAL":
        total_row = r
        break
assert total_row == 21, f"TOTAL attendu en 21, trouvé {total_row}"

heights = {
    r: ws.row_dimensions[r].height
    for r in list(ws.row_dimensions.keys())
    if ws.row_dimensions[r].height is not None
}

ws.insert_rows(total_row)
new_row = total_row  # 21
src_row = 20
for c in range(2, 21):
    src = ws.cell(row=src_row, column=c)
    dst = ws.cell(row=new_row, column=c)
    dst._style = copy(src._style)

for old_r, h in heights.items():
    target = old_r + 1 if old_r >= total_row else old_r
    ws.row_dimensions[target].height = h

values = {
    2: "TASK-017",
    3: "PRODUIT",
    4: "DevOS — Authentification (FEATURE-AUTH, Task 26)",
    5: "Durcissement sécurité pré-prod : rate limiting (DEBT-021) + CSP (DEBT-013) + REV-004 (Task 27)",
    6: (
        "Suite du plan d'action AUDIT-003 : src/lib/rate-limit.ts (fenêtre fixe mémoire, login 10/15 min/IP, "
        "register 5/h/IP, 429 uniforme { error } + Retry-After, 0 middleware/0 dépendance) ; CSP 9 directives "
        "dans next.config.ts (unsafe-inline/eval justifiés, frame-ancestors 'none', XFO SAMEORIGIN→DENY) ; "
        "revue indépendante REV-004 (Agent Reviewer) sur FEATURE-AUTH + fix seed + durcissement"
    ),
    7: "ORCHESTRATEUR (CTO) + AGENT REVIEWER (indépendant)",
    8: "TERMINÉ",
    9: 1,
    10: "P1",
    11: "Plan d'action AUDIT-003 (actions 1, 2, 4, 5), DEBT-013, DEBT-021",
    12: "src/lib/rate-limit.ts; src/lib/api-helpers.ts; src/app/api/auth/{login,register}/route.ts; next.config.ts; .ai/REVIEW_LOG.md (REV-004)",
    13: (
        "Saturation curl : login 10 autorisées (1 succès navigateur + 9×401) puis 429 Retry-After 877 s, "
        "register 5×400 puis 429, métier non affecté (tasks 401, health 200) · CSP : navigation complète "
        "AuthGate→connexion→3 onglets→déconnexion 0 violation console, HMR opérationnel · lint 0 · tsc 0 · "
        "REV-004 : 2 appels auth sous throttle (429 + Retry-After 643 = comportement correct)"
    ),
    14: "PASS",
    15: 0,
    16: "88f334f (rate limiting), db13a8b (CSP) — REV-004 : 0 code modifié (audit pur)",
    17: "REV-004 ✅ APPROUVÉE — 0 bloquant · 0 majeur · 4 mineurs tracés (XFF accès direct, Map > 5000, mdp démo, h1 AuthGate)",
    18: "2026-10-25",
    19: "2026-10-25",
    20: (
        "Dette 22 items : 17 TRAITÉ (013/021 soldés) · 2 EN COURS (003, 017) · 1 REPORTÉ (010) · 1 ACCEPTÉE "
        "(012) · 1 OUVERT (018) · restant avant exposition prod réelle : DEBT-020 purge comptes de démo + "
        "verrou accès direct :3000 (mineur REV-004)"
    ),
}
for c, v in values.items():
    ws.cell(row=new_row, column=c).value = v

tr = total_row + 1
ws.cell(row=tr, column=3).value = "17 tâches"
ws.cell(row=tr, column=9).value = 1
ws.cell(row=tr, column=15).value = 0
ws.cell(row=tr, column=20).value = "Terminées : 17 · En cours : 0 · À faire : 0 · Bloquées : 0"

st = wb["STATISTIQUES"]
st["C5"] = 17
st["C6"] = '=COUNTIF(TASKS!H5:H21,"TERMINÉ")'
st["C7"] = '=COUNTIF(TASKS!H5:H21,"EN COURS")'
st["C8"] = '=COUNTIF(TASKS!H5:H21,"À FAIRE")'
st["C9"] = '=COUNTIF(TASKS!H5:H21,"BLOQUÉ")'
st["C10"] = "=SUM(TASKS!O5:O21)"
st["C11"] = "=AVERAGE(TASKS!I5:I21)"
st["C12"] = "77 / 100 (AUDIT-003) · AUTORISÉ SOUS CONDITION — toutes conditions soldées + REV-004 APPROUVÉE"
st["D12"] = ".ai/AUDIT_REPORT.md"
st["C13"] = 22
st["D13"] = ".ai/DEBT_REPORT.md (17 TRAITÉ · 2 EN COURS · 1 REPORTÉ · 1 ACCEPTÉE · 1 OUVERT)"

wb.save(PATH)
print("TASKS.xlsx mis à jour : TASK-017 insérée en ligne 21, TOTAL en 22, STATISTIQUES 5:21 resynchronisées.")
