"""Task 29 — TASKS.xlsx : insertion TASK-018 (responsivité mobile vue de connexion) + resynchronisation.
Motif : copies de styles ligne 21 (TASK-017), plages STATISTIQUES étendues à 5:22, TOTAL 18 tâches.
"""
from copy import copy

from openpyxl import load_workbook

PATH = "/home/z/my-project/.ai/TASKS.xlsx"

wb = load_workbook(PATH)
ws = wb["TASKS"]

# 1. Localiser la ligne TOTAL
total_row = None
for r in range(5, ws.max_row + 1):
    if ws.cell(row=r, column=2).value == "TOTAL":
        total_row = r
        break
assert total_row == 22, f"TOTAL attendu en 22, trouvé {total_row}"

heights = {
    r: ws.row_dimensions[r].height
    for r in list(ws.row_dimensions.keys())
    if ws.row_dimensions[r].height is not None
}

# 2. Insérer la ligne TASK-018 juste avant TOTAL (styles copiés de TASK-017, ligne 21)
ws.insert_rows(total_row)
new_row = total_row  # 22
src_row = 21
for c in range(2, 21):
    src = ws.cell(row=src_row, column=c)
    dst = ws.cell(row=new_row, column=c)
    dst._style = copy(src._style)

# restaurer les hauteurs décalées
for old_r, h in heights.items():
    target = old_r + 1 if old_r >= total_row else old_r
    ws.row_dimensions[target].height = h

values = {
    2: "TASK-018",
    3: "PRODUIT",
    4: "DevOS — Authentification (FEATURE-AUTH)",
    5: "Responsivité mobile de la vue de connexion (AuthGate)",
    6: (
        "Exigence CTO « assure-toi de la responsivité mobile de la vue de connexion ». Analyse statique "
        "(auth-gate.tsx, input/card/tabs/toast, layout) + tests navigateur multi-viewports (375 / 320 / "
        "667×375 paysage / 1280). Correctifs : (1) min-h-dvh + centrage « sûr » my-auto (remplace "
        "justify-center : en hauteur réduite — clavier iOS, paysage — le haut du formulaire reste "
        "atteignable au scroll, plus de contenu clippé) ; (2) paddings scoped px-4 sm:px-6 sur "
        "CardHeader/CardContent + gap-4/py-4 sm:gap-6/sm:py-6 (rythme 320 px) ; (3) h1 sr-only (partie "
        "h1 de DEBT-017) ; écran de chargement aligné en min-h-dvh. Aucun changement de comportement "
        "fonctionnel ni de contrat API."
    ),
    7: "Frontend / UX (CTO Task 29)",
    8: "TERMINÉ",
    9: 1,
    10: "P2",
    11: "TASK-015 (FEATURE-AUTH — la vue de connexion existe)",
    12: (
        "src/components/assistant/auth-gate.tsx; tool-results/task29-mobile-375-before.png; "
        "tool-results/task29-mobile-320-before.png; tool-results/task29-mobile-375-after.png; "
        "tool-results/task29-mobile-320-after.png; tool-results/task29-mobile-landscape-after.png; "
        "tool-results/task29-desktop-after.png; tool-results/task29-mobile-login-ok.png; worklog.md"
    ),
    13: (
        "hasHScroll=false à 320/375/667/1280 · paysage 667×375 : docH 606 > viewport 375, scroll-top → "
        "cardTop=40 px (haut du formulaire visible — centrage sûr prouvé) · computed minHeight=100dvh · "
        "h1 sr-only présent (position absolute) · E2E mobile complet : connexion demo@devos.test → "
        "« Bienvenue, Démonstration DevOS ! » + app shell (Assistant/Tâches/Paramètres) · 0 erreur console "
        "· toast mobile top full-width observé live"
    ),
    14: "PASS",
    15: 0,
    16: (
        "0 middleware respecté (ADR-003) · aucun composant ui/ global modifié (overrides scoped à "
        "l'instance AuthForm) · dérive environnement détectée et corrigée en passant : .env sans "
        "JWT_SECRET (restauration env) → secret régénéré + User table vide → seed rejoué (comptes "
        "auto-sauvegarde, cf. COMMIT_LOG §auto-sauvegarde)"
    ),
    17: (
        "Vue de connexion : 100 % utilisable 320→1280 px, clavier iOS et paysage couverts (centrage sûr) · "
        "a11y : h1 ajouté (partie de DEBT-017 — le découpage UI reste EN COURS)"
    ),
    18: "2026-10-26",
    19: "2026-10-26",
    20: (
        "min-h-dvh : dégradation gracieuse top-alignée sur navigateurs sans dvh (< 2022) — acceptable · "
        "AUDIT-004 requis après cette évolution (règle « audit après chaque feature »)"
    ),
}
for c, v in values.items():
    ws.cell(row=new_row, column=c).value = v

# 3. Ligne TOTAL (décalée en 23)
tr = total_row + 1
ws.cell(row=tr, column=3).value = "18 tâches"
ws.cell(row=tr, column=9).value = 1
ws.cell(row=tr, column=15).value = 0
ws.cell(row=tr, column=20).value = "Terminées : 18 · En cours : 0 · À faire : 0 · Bloquées : 0"

# 4. STATISTIQUES — plages étendues aux lignes 5..22 (18 tâches)
st = wb["STATISTIQUES"]
st["C5"] = 18
st["C6"] = '=COUNTIF(TASKS!H5:H22,"TERMINÉ")'
st["C7"] = '=COUNTIF(TASKS!H5:H22,"EN COURS")'
st["C8"] = '=COUNTIF(TASKS!H5:H22,"À FAIRE")'
st["C9"] = '=COUNTIF(TASKS!H5:H22,"BLOQUÉ")'
st["C10"] = "=SUM(TASKS!O5:O22)"
st["C11"] = "=AVERAGE(TASKS!I5:I22)"

wb.save(PATH)
print("TASKS.xlsx mis à jour : TASK-018 insérée en ligne 22, TOTAL en 23, STATISTIQUES resynchronisées (5:22).")
