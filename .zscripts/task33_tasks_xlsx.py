"""Task 33 — TASKS.xlsx : insertion TASK-019 (FEATURE-UI-SHELL) + resynchronisation.
Motif : copies de styles ligne 22 (TASK-018), plages STATISTIQUES étendues à 5:23, TOTAL 19 tâches.
"""
from copy import copy

from openpyxl import load_workbook

PATH = "/home/z/my-project/.ai/TASKS.xlsx"

wb = load_workbook(PATH)
ws = wb["TASKS"]

total_row = None
for r in range(5, ws.max_row + 1):
    if ws.cell(row=r, column=2).value == "TOTAL":
        total_row = r
        break
assert total_row == 23, f"TOTAL attendu en 23, trouvé {total_row}"

heights = {
    r: ws.row_dimensions[r].height
    for r in list(ws.row_dimensions.keys())
    if ws.row_dimensions[r].height is not None
}

ws.insert_rows(total_row)
new_row = total_row  # 23
src_row = 22
for c in range(2, 21):
    src = ws.cell(row=src_row, column=c)
    dst = ws.cell(row=new_row, column=c)
    dst._style = copy(src._style)

for old_r, h in heights.items():
    target = old_r + 1 if old_r >= total_row else old_r
    ws.row_dimensions[target].height = h

values = {
    2: "TASK-019",
    3: "PRODUIT",
    4: "DevOS — Espace de travail (FEATURE-UI-SHELL, Task 33)",
    5: "Shell complet : sidebar 4 sections, permissions centralisées, 12 vues fonctionnelles, navigation mobile 5 destinations",
    6: (
        "Exigence CTO 2026-10-26 (cahier des charges interface, recommandation finale §6 : sidebar stable "
        "+ permissions centralisées + vues réellement fonctionnelles d'abord). Livré : registre de navigation "
        "central (lib/navigation.ts, 17 destinations, statut active/planned) + permissions centralisées "
        "(lib/permissions.ts, matrice ROLES_PERMISSIONS.md, canAccessView) ; sidebar desktop + tiroir mobile "
        "(Sheet) + bottom nav 5 destinations (Accueil/Projets/Tâches/Assistant/Plus) ; Prisma +2 modèles "
        "(Project, Note) + Task.projectId (db:push) ; 6 endpoints v1.3 (projects CRUD, notes CRUD, activity, "
        "ai/chat LLM rate-limité 10/5min) ; 12 vues : Tableau de bord (données réelles), Projets (favoris/"
        "archivage/progression), Tâches (7 smart views + form dialog + TaskCard, TasksManager 531 l. supprimé), "
        "Planning (calendrier mensuel), Documents & Notes (recherche/favoris/corbeille), Assistant IA (SDK "
        "serveur, 3 contextes), Assistant vocal (AssistantPanel réutilisé), Journal d'activité (ReminderLog "
        "réels, 4 onglets), Notifications, Synchronisation, Paramètres (SettingsScreen réutilisé), ComingSoon "
        "générique pour 5 menus avancés (badge « Bientôt » — ADR-004) ; ADR-004 + ROLES_PERMISSIONS.md"
    ),
    7: "Équipe complète (CTO orchestrateur)",
    8: "TERMINÉ",
    9: 1,
    10: "P1",
    11: "TASK-015 (FEATURE-AUTH), TASK-018 (responsive), exigence CTO cahier des charges interface",
    12: (
        "src/lib/navigation.ts; src/lib/permissions.ts; prisma/schema.prisma; src/app/api/{projects,notes,"
        "activity,ai}; src/components/shell/; src/components/views/; src/components/assistant/dev-os-app.tsx; "
        ".ai/ADR/ADR-004-shell-navigation-permissions.md; .ai/ROLES_PERMISSIONS.md; .ai/API_CONTRACTS.md v1.3; "
        "tool-results/task33-*.png"
    ),
    13: (
        "bun run lint exit 0 · bunx tsc exit 0 · curl API : projects 200/201, notes 201, tasks+projectId 201 "
        "(projectName résolu), activity 200, ai/chat LLM réel en fr (remaining décrémenté) · E2E navigateur : "
        "dashboard (compteurs réels + scheduler actif), projets (création UI + progression), tâches (smart "
        "views + dialog), planning (calendrier), IA (conversation complète en fr), journal (entrées réelles), "
        "mobile 375 (bottom nav 5 destinations + tiroir complet + vue Bientôt), 0 erreur console · purge "
        "artefacts E2E (DEBT-020)"
    ),
    14: "PASS",
    15: 0,
    16: (
        "0 middleware respecté (ADR-003) · API-First : contrat v1.3 écrit AVANT l'UI · IA sans accès outillé "
        "aux données (ROLES_PERMISSIONS §3) · fichiers < 500 l. sauf SettingsScreen réutilisé (632, DEBT-017 "
        "continue) · correctifs en vol : alias filtré Prisma _count → groupBy, bouton Plus toujours présent, "
        "garde-fou navigate() retiré (vues planned → ComingSoonView)"
    ),
    17: (
        "Espace de travail MVP conforme au cahier des charges : 12 vues actives + 5 destinations « Bientôt » "
        "visibles · navigation adaptative desktop/mobile · permissions centralisées prêtes pour le "
        "multi-rôles · backlogs tracés : historique IA persistant, mode mains libres, automatisations, Git/"
        "GitHub"
    ),
    18: "2026-10-26",
    19: "2026-10-26",
    20: (
        "Périmètre volontairement MVP (recommandation CTO §6) · AUDIT-005 requis après cette feature · "
        "DEBT-017 : settings-screen.tsx (632) reste à découper"
    ),
}
for c, v in values.items():
    ws.cell(row=new_row, column=c).value = v

tr = total_row + 1
ws.cell(row=tr, column=3).value = "19 tâches"
ws.cell(row=tr, column=9).value = 1
ws.cell(row=tr, column=15).value = 0
ws.cell(row=tr, column=20).value = "Terminées : 19 · En cours : 0 · À faire : 0 · Bloquées : 0"

st = wb["STATISTIQUES"]
st["C5"] = 19
st["C6"] = '=COUNTIF(TASKS!H5:H23,"TERMINÉ")'
st["C7"] = '=COUNTIF(TASKS!H5:H23,"EN COURS")'
st["C8"] = '=COUNTIF(TASKS!H5:H23,"À FAIRE")'
st["C9"] = '=COUNTIF(TASKS!H5:H23,"BLOQUÉ")'
st["C10"] = "=SUM(TASKS!O5:O23)"
st["C11"] = "=AVERAGE(TASKS!I5:I23)"

wb.save(PATH)
print("TASKS.xlsx : TASK-019 insérée en ligne 23, TOTAL en 24, STATISTIQUES 5:23.")
